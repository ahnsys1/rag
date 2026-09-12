"""Extract plain text from uploaded files (PDF, DOCX, TXT/MD)."""

import io
import re
from dataclasses import dataclass, field
from pathlib import PurePosixPath

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Page:
    number: int | None
    text: str


@dataclass
class ExtractedDocument:
    pages: list[Page] = field(default_factory=list)

    @property
    def page_count(self) -> int | None:
        numbered = [p for p in self.pages if p.number is not None]
        return len(numbered) or None


_WS_RE = re.compile(r"[ \t\u00a0]+")
_NL_RE = re.compile(r"\n{3,}")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _NL_RE.sub("\n\n", text).strip()


def extract_pdf(data: bytes) -> ExtractedDocument:
    reader = PdfReader(io.BytesIO(data))
    doc = ExtractedDocument()
    for i, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if text:
            doc.pages.append(Page(number=i, text=text))
    return doc


def extract_docx(data: bytes) -> ExtractedDocument:
    document = DocxDocument(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = normalize_text("\n".join(parts))
    return ExtractedDocument(pages=[Page(number=None, text=text)] if text else [])


def extract_plain(data: bytes) -> ExtractedDocument:
    text = normalize_text(data.decode("utf-8", errors="replace"))
    return ExtractedDocument(pages=[Page(number=None, text=text)] if text else [])


def extract(filename: str, data: bytes) -> ExtractedDocument:
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext in {".txt", ".md"}:
        return extract_plain(data)
    raise UnsupportedFileType(
        f"Unsupported file type '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )
