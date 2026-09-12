from app.ingestion.chunking import split_text
from app.ingestion.loaders import ExtractedDocument, Page, extract, extract_docx, normalize_text


def test_split_text_respects_size_and_overlap():
    text = " ".join(f"word{i}" for i in range(500))
    chunks = split_text(text, chunk_size=200, overlap=40)
    assert len(chunks) > 1
    assert all(len(c) <= 200 for c in chunks)
    # Overlap: last words of chunk N should appear at the start of chunk N+1.
    for prev, nxt in zip(chunks, chunks[1:], strict=False):
        assert prev.split()[-1] in nxt.split()[:8]


def test_split_text_keeps_paragraphs_together_when_small():
    text = "Prvni odstavec.\n\nDruhy odstavec."
    assert split_text(text, chunk_size=100, overlap=10) == [text]


def test_split_text_empty():
    assert split_text("   \n\n ", 100, 10) == []


def test_normalize_text_collapses_whitespace():
    assert normalize_text("a \t b\r\n\n\n\nc") == "a b\n\nc"


def test_extract_plain_and_page_count():
    doc = extract("notes.txt", b"hello world")
    assert doc.pages[0].text == "hello world"
    assert doc.page_count is None
    assert ExtractedDocument(pages=[Page(1, "a"), Page(2, "b")]).page_count == 2


def test_extract_rejects_unknown_extension():
    import pytest

    from app.ingestion.loaders import UnsupportedFileType

    with pytest.raises(UnsupportedFileType):
        extract("archive.zip", b"PK")


def test_extract_docx_roundtrip():
    import io

    from docx import Document

    buf = io.BytesIO()
    d = Document()
    d.add_paragraph("Ahoj světe")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    d.save(buf)

    out = extract_docx(buf.getvalue())
    assert "Ahoj světe" in out.pages[0].text
    assert "A | B" in out.pages[0].text
